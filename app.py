import streamlit as st
from docx import Document
from openpyxl import load_workbook
import unicodedata


# =========================
# NORMALIZAÇÃO
# =========================
def normalizar(texto):
    if not texto:
        return ""
    texto = str(texto).strip().lower()
    texto = ''.join(
        c for c in unicodedata.normalize('NFD', texto)
        if unicodedata.category(c) != 'Mn'
    )
    return texto


def limpar_item_excel(texto):
    if not texto:
        return ""
    texto = str(texto).split("&")[0]
    return texto.strip()


# =========================
# DOCX
# =========================
@st.cache_data
def ler_docx(file):
    doc = Document(file)
    return [
        row.cells[0].text.strip()
        for row in doc.tables[0].rows
        if row.cells[0].text.strip()
    ]


# =========================
# EXCEL
# =========================
@st.cache_data
def analisar_excel(file, substancias):
    wb = load_workbook(file, read_only=True, data_only=True)
    sheets = wb.worksheets[:2]

    substancias_norm = {normalizar(s): s for s in substancias}
    resultado = {s: set() for s in substancias}

    for sheet in sheets:
        for row in sheet.iter_rows(min_row=6, values_only=True):
            if not row:
                continue

            if len(row) < 7:
                continue

            area = row[6]
            if not area:
                continue

            area = str(area).strip()

            # construir set da linha (MATCH RÁPIDO E EXACTO)
            row_set = set()

            for cell in row:
                if cell:
                    row_set.add(normalizar(cell))

            # matching exato
            for norm_sub, original_sub in substancias_norm.items():
                if norm_sub in row_set:
                    resultado[original_sub].add(area)

    output = [
        (s, sorted(a))
        for s, a in resultado.items()
        if a
    ]

    wb.close()
    return output


# =========================
# STREAMLIT UI
# =========================
st.title("🔬 Analisador de Substâncias")

docx_file = st.file_uploader("DOCX", type=["docx"])
xlsx_file = st.file_uploader("XLSX", type=["xlsx"])

if docx_file and xlsx_file:

    if st.button("Analisar"):

        st.info("A processar DOCX...")
        substancias = ler_docx(docx_file)

        st.info(f"Substâncias: {len(substancias)}")

        st.info("A analisar Excel...")

        with st.spinner("Processando..."):
            resultado = analisar_excel(xlsx_file, substancias)

        st.success("Concluído!")

        for substancia, areas in resultado:
            st.write(f"**{substancia}** → {', '.join(areas)}")

